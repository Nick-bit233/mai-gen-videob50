"""Read custom charts from XLSX with explicit column mapping and row diagnostics."""
from dataclasses import dataclass
from decimal import Decimal
from difflib import SequenceMatcher
import io
import math
from pathlib import Path
import re
import unicodedata
from uuid import uuid4
import zipfile

from openpyxl import load_workbook
from PIL import Image

from utils.ChartUtils import validate_complete_record_form, get_fc_status_options, get_fs_status_options
from utils.ChartUtils import try_parse_difficulty
from utils.DataUtils import compute_rating, compute_chunithm_rating

ALIASES = {
    'song_name': ('曲名', '歌曲名称', '歌曲名', 'title', 'song', 'song title'),
    'artist': ('曲师', '作曲家', '作曲', '作者', 'composer'),
    'song_id': ('歌曲ID', '曲目ID', '来源ID'),
    'chart_type': ('谱面类型', '类型', 'SD/DX', 'chart type', 'type'),
    'level_index': ('难度', '难度等级', '谱面难度', 'level label', 'level'),
    'difficulty': ('定数', '谱面定数', 'ds', 'constant'),
    'max_dx_score': ('最大DX分数', '理论DX分数', 'max dx score'),
    'achievement': ('达成率', '分数', '成绩', 'score', 'accuracy', 'achievement percentage'),
    'fc_status': ('FC', 'FC状态', '全连状态', 'combo'),
    'fs_status': ('FS', 'FS状态', '同步状态', 'full chain'),
    'play_count': ('游玩次数', '次数', 'play count'),
    'dx_score': ('DX分数', 'DX Score'), 'dx_rating': ('DX Rating', '舞萌Rating'),
    'chuni_rating': ('Chuni Rating', '中二Rating'),
    'clip_title_name': ('片段标题', '标题', '自定义标题', 'clip title'),
    'order_in_archive': ('顺序', '排序', 'order'),
    'jacket_path': ('曲绘路径', '曲绘文件', '曲绘', '图片路径', 'jacket', 'cover path'),
    'video_path': ('本地视频路径', '视频路径'),
    'video_source': ('视频链接', 'BV', 'BV号', 'video url', 'source url'),
    'source_pid': ('分P', '分 P', '视频分P', 'pid'),
}


def normalize(value):
    return re.sub(r'[\s_:\-/（）()]+', '', unicodedata.normalize('NFKC', str(value or '')).casefold())


def suggest_mapping(headers, game_type):
    aliases = {field: {normalize(field), *(normalize(v) for v in values)} for field, values in ALIASES.items()}
    aliases['chuni_rating' if game_type == 'chunithm' else 'dx_rating'].add('rating')
    mapping, suggestions = {}, {}
    for index, header in enumerate(headers):
        value = normalize(header)
        matches = [field for field, names in aliases.items() if value and value in names]
        scores = sorted(((max(SequenceMatcher(None, value, name).ratio() for name in names), field)
                         for field, names in aliases.items()), reverse=True) if value else []
        field = matches[0] if len(matches) == 1 else None
        # Fuzzy matching is a suggestion only; ambiguous/fuzzy columns require user mapping.
        mapping[index] = field
        suggestions[index] = matches or [field for score, field in scores[:3] if score >= .65]
    counts = {field: list(mapping.values()).count(field) for field in set(mapping.values()) if field}
    for index, field in list(mapping.items()):
        if field and counts[field] > 1:
            mapping[index] = None
    return mapping, suggestions


@dataclass
class ExcelCell:
    value: object
    number_format: str = 'General'
    uncached_formula: bool = False


@dataclass
class ExcelTable:
    headers: list[str]
    rows: list[tuple[int, list[ExcelCell]]]
    sheet_names: list[str]
    sheet: str


def read_excel(data: bytes, sheet=None):
    if not data or len(data) > 10 * 1024 * 1024:
        raise ValueError('Excel 文件为空或超过 10 MiB。')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(file.file_size for file in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError('Excel 解压后的内容过大。')
        if any('vbaProject' in file.filename for file in archive.infolist()):
            raise ValueError('请另存为无宏的 .xlsx。')
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        sheet = sheet or ('Sheet1' if 'Sheet1' in workbook.sheetnames else workbook.sheetnames[0])
        if sheet not in workbook.sheetnames:
            raise ValueError('选择的工作表不存在。')
        ws, values = workbook[sheet], cached[sheet]
        if (ws.max_row or 0) > 5001 or (ws.max_column or 0) > 64:
            raise ValueError('首期支持最多 5000 条记录和 64 列，请移除多余空白区域。')
        row_pairs = zip(ws.iter_rows(), values.iter_rows())
        first = next(row_pairs, None)
        if not first:
            raise ValueError('工作表为空。')
        headers = [str(cell.value or '').strip() for cell in first[0]]
        if len(headers) > 64:
            raise ValueError('首期支持最多 64 列。')
        if not any(headers):
            raise ValueError('首行需要填写表头。')
        rows = []
        for number, (raw, cached_row) in enumerate(row_pairs, start=2):
            if number > 5001 or len(raw) > 64:
                raise ValueError('首期支持最多 5000 条记录和 64 列。')
            if not any(cell.value is not None for cell in raw):
                continue
            cells = [ExcelCell(cache.value if cell.data_type == 'f' else cell.value,
                               cell.number_format, cell.data_type == 'f' and cache.value is None)
                     for cell, cache in zip(raw, cached_row)]
            rows.append((number, cells))
        return ExcelTable(headers, rows, list(workbook.sheetnames), sheet)
    finally:
        workbook.close()
        cached.close()


def _text(value):
    if value is None or isinstance(value, float) and not math.isfinite(value):
        return ''
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()


def local_image_path(value, base_dir=None):
    value = _text(value).strip('"')
    if not value:
        return None
    if '://' in value:
        raise ValueError('曲绘列需要本地文件路径。')
    path = Path(value)
    if not path.is_absolute():
        if not base_dir:
            raise ValueError('相对曲绘路径需要填写曲绘根目录。')
        base = Path(base_dir).resolve()
        path = (base / path).resolve()
        if not path.is_relative_to(base):
            raise ValueError('相对曲绘路径超出选定根目录。')
    if not path.is_file():
        raise ValueError('曲绘文件不存在。')
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError('曲绘文件超过 20 MiB。')
    with Image.open(path) as image:
        if image.width * image.height > 40_000_000:
            raise ValueError('曲绘像素尺寸过大。')
        image.verify()
    return str(path.resolve())


def convert_rows(table, mapping, game_type, *, jacket_root=None, sort_by_order=False):
    if game_type not in {'maimai', 'chunithm'}:
        raise ValueError('请选择舞萌或中二。')
    fields = [field for field in mapping.values() if field]
    if any(field not in ALIASES for field in fields) or len(fields) != len(set(fields)):
        raise ValueError('每个字段只能对应一列，请修正重复列映射。')
    batch = uuid4().hex
    records, preview = [], []
    for row_number, cells in table.rows:
        warnings = []
        raw = {}
        for index, field in mapping.items():
            if field and index < len(cells):
                cell = cells[index]
                value = cell.value
                if cell.uncached_formula:
                    warnings.append(f'{field} 公式无缓存，请保存数值；本次使用缺失默认值')
                if field == 'achievement' and game_type == 'maimai':
                    if isinstance(value, (int, float)) and '%' in cell.number_format:
                        value = str(Decimal(str(value)) * 100)
                    elif isinstance(value, str) and value.strip().endswith('%'):
                        value = value.strip()[:-1]
                raw[field] = _text(value)
        values = {field: raw.get(field, '') for field in ALIASES}
        for field, fallback in {'song_name': f'未命名谱面 第{row_number}行', 'artist': '未知曲师',
                                 'difficulty': '?', 'clip_title_name': f'Clip_{len(records)+1}'}.items():
            if not values[field]:
                values[field] = fallback
                warnings.append(f'{field} 缺失，使用 {fallback}')
        values['song_id'] = f'custom_excel_{batch}_{row_number}'
        level_names = ['BASIC', 'ADVANCED', 'EXPERT', 'MASTER', 'RE:MASTER' if game_type == 'maimai' else 'ULTIMA']
        level_aliases = {normalize(v): i for i, v in enumerate(level_names)}
        level_aliases.update({v: i for i, v in enumerate(('绿', '黄', '红', '紫', '白' if game_type == 'maimai' else '黑'))})
        for field, aliases, default in (
            ('level_index', level_aliases, 3),
            ('chart_type', {'sd': 0, 'standard': 0, '标准': 0, 'normal': 0, '普通': 0, **({'dx': 1, 'deluxe': 1} if game_type == 'maimai' else {})}, 0),
        ):
            text = normalize(values[field])
            allowed = range(5) if field == 'level_index' else range(2 if game_type == 'maimai' else 1)
            converted = int(text) if text.isdigit() and int(text) in allowed else aliases.get(text)
            if converted is None:
                converted = default
                if text:
                    warnings.append(f'{field} 无法识别，使用 {default}')
            values[field] = converted
        for field, options in [('fc_status', get_fc_status_options(game_type)), ('fs_status', get_fs_status_options(game_type))]:
            text = normalize(values[field]).replace('＋', '+')
            text = {'fc+': 'fcp', 'ap+': 'app', 'fs+': 'fsp', 'fsd+': 'fsdp', 'aj+': 'ajc', '无': 'none', '': 'none'}.get(text, text)
            if text not in options:
                warnings.append(f'{field} 无法识别，使用 none')
                text = 'none'
            values[field] = text
        for field in ('achievement', 'max_dx_score', 'dx_score', 'dx_rating', 'chuni_rating', 'play_count'):
            text = values[field] or '0'
            try:
                value = float(text)
                integer = field in {'max_dx_score', 'dx_score', 'dx_rating', 'play_count'} or field == 'achievement' and game_type == 'chunithm'
                if not math.isfinite(value) or value < 0 or integer and (not value.is_integer() or value > 2**63-1) or field == 'achievement' and value > (101 if game_type == 'maimai' else 1010000):
                    raise ValueError()
            except (ValueError, OverflowError):
                warnings.append(f'{field} 数值无效，使用 0；可返回修改表格')
                value = 0
            values[field] = str(int(value)) if integer else str(value)
        rating_key = 'dx_rating' if game_type == 'maimai' else 'chuni_rating'
        if not raw.get(rating_key):
            ds = try_parse_difficulty(values['difficulty'])
            ds = ds if ds is not None and math.isfinite(ds) else 0
            values[rating_key] = str(compute_rating(ds, float(values['achievement'])) if game_type == 'maimai' else compute_chunithm_rating(ds, int(float(values['achievement']))))
        record, errors = validate_complete_record_form(values, game_type, auto_calc_rating=False)
        if errors:
            preview.append({'Excel行': row_number, '曲名': values['song_name'], '状态': '无法导入', '提示': '；'.join(errors)})
            continue
        record['chart_data']['game_type'] = game_type
        record['raw_data'] = {'excel_row': row_number, 'source_song_id': raw.get('song_id', ''), 'rating_explicit': bool(raw.get(rating_key))}
        if values['jacket_path']:
            try:
                record['_import_jacket_path'] = local_image_path(values['jacket_path'], jacket_root)
            except (OSError, ValueError) as exc:
                warnings.append(f'曲绘留空：{exc}')
        if values['video_path']:
            path = Path(values['video_path'])
            if path.is_file():
                record['chart_data']['video_path'] = str(path.resolve())
            else:
                warnings.append('本地视频不存在，路径留空')
        if values['video_source']:
            try:
                from utils.video_metadata import parse_bilibili_reference
                record['chart_data']['video_metadata'] = parse_bilibili_reference(values['video_source'], title=values['song_name'], requested_pid=int(values['source_pid']) if values['source_pid'] else None)
            except ValueError as exc:
                warnings.append(f'视频参考留空：{exc}')
        record['_excel_order'] = len(records)
        if sort_by_order and values['order_in_archive']:
            try:
                value = float(values['order_in_archive'])
                if not math.isfinite(value):
                    raise ValueError()
                record['_excel_order'] = value
            except ValueError:
                warnings.append('排序值无效，保留行顺序')
        records.append(record)
        preview.append({'Excel行': row_number, '曲名': values['song_name'], '曲师': values['artist'],
                        '定数': values['difficulty'], '成绩': record['achievement'], 'Rating': record[rating_key],
                        '曲绘': record.get('_import_jacket_path', ''), '状态': '默认值待确认' if warnings else '可导入',
                        '提示': '；'.join(warnings)})
    if sort_by_order:
        records.sort(key=lambda r: r['_excel_order'])
    for index, record in enumerate(records):
        record['order_in_archive'] = index
        record.pop('_excel_order', None)
    return records, preview


def metadata_candidates(record, songs):
    chart = record['chart_data']
    game = chart['game_type']
    title = normalize(chart['song_name'])
    artist = normalize(chart['artist'])
    candidates = []
    for song in songs:
        similarity = SequenceMatcher(None, title, normalize(song.get('title'))).ratio()
        if similarity < .8 or artist not in {'', normalize('未知曲师')} and artist != normalize(song.get('artist')):
            continue
        for candidate in song.get('charts_info', []):
            expected_type = ('dx' if chart['chart_type'] == 1 else 'standard') if game == 'maimai' else 'normal'
            if candidate.get('difficulty') != chart['level_index'] or candidate.get('type', 'normal') not in {expected_type, 'standard' if game == 'chunithm' else expected_type}:
                continue
            ds = candidate.get('level_value_cn')
            if ds is None:
                ds = candidate.get('level_value_latest')
            if ds is None or not isinstance(ds, (int,float)) or not math.isfinite(ds):
                continue
            candidates.append({'title': song['title'], 'artist': song.get('artist', ''), 'difficulty': str(ds), 'similarity': similarity})
    return sorted(candidates, key=lambda c: c['similarity'], reverse=True)[:5]


def apply_metadata_candidate(record, candidate):
    chart = record['chart_data']
    if chart['artist'] in {'', '未知曲师'}:
        chart['artist'] = candidate['artist']
    current = try_parse_difficulty(chart['difficulty'])
    if current is None or not math.isfinite(current):
        chart['difficulty'] = candidate['difficulty']
    if not record.get('raw_data', {}).get('rating_explicit'):
        ds = try_parse_difficulty(chart['difficulty']) or 0
        game = chart['game_type']
        rating_key = 'dx_rating' if game == 'maimai' else 'chuni_rating'
        record[rating_key] = compute_rating(ds, record['achievement']) if game == 'maimai' else compute_chunithm_rating(ds, record['achievement'])
